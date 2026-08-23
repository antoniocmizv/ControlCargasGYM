"""Hojas por sesión: la parrilla que sustituye a la plantilla de papel."""

from datetime import date, timedelta
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.conftest import auth_headers

HOY = date.today()
# Día propio de estas pruebas: otras crean sesiones para hoy y contaminarían la hoja.
DIA = HOY - timedelta(days=200)


@pytest.fixture
def sesion_con_cargas(client: TestClient, coach_headers: dict) -> dict:
    jugadores = [
        client.post(
            "/api/coach/players", json={"name": nombre, "pin": pin}, headers=coach_headers
        ).json()
        for nombre, pin in [("Grid Ana", "1111"), ("Grid Luis", "2222"), ("Grid Marta", "3333")]
    ]
    ejercicios = client.get("/api/coach/exercises", headers=coach_headers).json()

    routine = client.post(
        "/api/coach/routines",
        headers=coach_headers,
        json={
            "name": "Parrilla",
            "session_date": str(DIA),
            "items": [
                {"exercise_id": ejercicios[0]["id"], "sets": 3, "target_reps": 8},
                {"exercise_id": ejercicios[1]["id"], "sets": 2, "target_reps": 10},
            ],
            "assignments": [
                {"target_type": "player", "user_id": jugador["id"]} for jugador in jugadores
            ],
        },
    ).json()

    # Ana completa el primer ejercicio; Marta deja la última serie sin registrar.
    cargas = {0: [(1, 80), (2, 85), (3, 90)], 2: [(1, 70), (2, 72.5)]}
    for indice, (jugador, pin) in enumerate(zip(jugadores, ["1111", "2222", "3333"])):
        headers = auth_headers(
            client.post(
                "/api/auth/login/player", json={"user_id": jugador["id"], "pin": pin}
            ).json()["access_token"]
        )
        for numero, kg in cargas.get(indice, []):
            client.post(
                "/api/logs",
                headers=headers,
                json={
                    "routine_exercise_id": routine["items"][0]["id"],
                    "set_number": numero,
                    "load_kg": kg,
                    "reps": 8,
                },
            )

    return {"routine": routine, "jugadores": jugadores}


def _descarga(client: TestClient, coach_headers: dict, **params):
    respuesta = client.get("/api/export/excel", params=params, headers=coach_headers)
    assert respuesta.status_code == 200
    return load_workbook(BytesIO(respuesta.content))


def test_sin_pedirlo_el_excel_no_cambia(client: TestClient, coach_headers: dict, sesion_con_cargas):
    libro = _descarga(client, coach_headers)
    assert libro.sheetnames == ["Cargas", "Resumen"]


def test_la_hoja_de_sesion_lleva_fecha_y_nombre(
    client: TestClient, coach_headers: dict, sesion_con_cargas
):
    libro = _descarga(client, coach_headers, por_sesion="true", date_from=str(DIA), date_to=str(DIA))

    assert "Cargas" in libro.sheetnames and "Resumen" in libro.sheetnames
    hoja = next(n for n in libro.sheetnames if n.startswith(f"{DIA:%Y-%m-%d}"))
    assert "Parrilla" in hoja
    assert len(hoja) <= 31, "Excel no admite nombres de hoja de más de 31 caracteres"


def test_la_parrilla_coloca_jugadores_en_filas_y_series_en_columnas(
    client: TestClient, coach_headers: dict, sesion_con_cargas
):
    libro = _descarga(client, coach_headers, por_sesion="true", date_from=str(DIA), date_to=str(DIA))
    hoja = libro[next(n for n in libro.sheetnames if n.startswith(f"{DIA:%Y-%m-%d}"))]

    assert hoja.cell(row=5, column=1).value == "Jugador"
    assert [hoja.cell(row=5, column=c).value for c in range(2, 7)] == ["S1", "S2", "S3", "S1", "S2"]

    filas = {hoja.cell(row=f, column=1).value: f for f in range(6, 9)}
    assert set(filas) == {"Grid Ana", "Grid Luis", "Grid Marta"}

    ana = [hoja.cell(row=filas["Grid Ana"], column=c).value for c in range(2, 5)]
    assert ana == [80, 85, 90]


def test_las_series_sin_registrar_salen_marcadas(
    client: TestClient, coach_headers: dict, sesion_con_cargas
):
    libro = _descarga(client, coach_headers, por_sesion="true", date_from=str(DIA), date_to=str(DIA))
    hoja = libro[next(n for n in libro.sheetnames if n.startswith(f"{DIA:%Y-%m-%d}"))]

    filas = {hoja.cell(row=f, column=1).value: f for f in range(6, 9)}
    marta = [hoja.cell(row=filas["Grid Marta"], column=c).value for c in range(2, 5)]
    assert marta == [70, 72.5, "—"], "la tercera serie de Marta está sin registrar"

    luis = [hoja.cell(row=filas["Grid Luis"], column=c).value for c in range(2, 5)]
    assert luis == ["—", "—", "—"], "Luis no registró nada"


def test_filtrar_por_jugador_deja_solo_su_fila(
    client: TestClient, coach_headers: dict, sesion_con_cargas
):
    ana = next(j for j in sesion_con_cargas["jugadores"] if j["name"] == "Grid Ana")
    libro = _descarga(
        client,
        coach_headers,
        por_sesion="true",
        player_id=ana["id"],
        date_from=str(DIA),
        date_to=str(DIA),
    )
    hoja = libro[next(n for n in libro.sheetnames if n.startswith(f"{DIA:%Y-%m-%d}"))]

    assert hoja.cell(row=6, column=1).value == "Grid Ana"
    assert hoja.cell(row=7, column=1).value is None


def test_dos_sesiones_el_mismo_dia_no_chocan_de_nombre(
    client: TestClient, coach_headers: dict
):
    ejercicio = client.get("/api/coach/exercises", headers=coach_headers).json()[0]
    dia = HOY - timedelta(days=40)
    for _ in range(2):
        client.post(
            "/api/coach/routines",
            headers=coach_headers,
            json={
                "name": "Repetida",
                "session_date": str(dia),
                "items": [{"exercise_id": ejercicio["id"], "sets": 1}],
                "assignments": [{"target_type": "all"}],
            },
        )

    libro = _descarga(client, coach_headers, por_sesion="true", date_from=str(dia), date_to=str(dia))
    hojas = [n for n in libro.sheetnames if n.startswith(f"{dia:%Y-%m-%d}")]
    assert len(hojas) == 2
    assert len(set(hojas)) == 2, "los nombres de hoja deben ser únicos"
